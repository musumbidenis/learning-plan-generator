"""Module D3 - assessment_docs tests. ZERO API calls, no sample PDFs.

Every fixture is hand-made here so the numbers are known: the four documents are
built from it and asserted on as produced .docx (no golden files), the way
tests/test_doc_builder.py does.
"""

import io

from docx import Document

import assessment_docs
from assessment_models import (ANALYSING, APPLYING, CAT_1, CAT_2, KNOWLEDGE,
                               PRACTICAL, THEORY, UNDERSTANDING, Allocation,
                               AssessmentTool, CatDefinition, ChecklistItem,
                               Item, MarkingPoint, OralQuestion, Scenario,
                               TaskBrief, UnitWeighting, WeightedElement,
                               WeightedPC)

UNIT_TITLE = "APPLY COMPUTER PROGRAMMING PRINCIPLES"
ISCED = "0613 451 05A"
CDACC = "IT/CU/ICTA/CC/02/5/MA"


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
def _weighting() -> UnitWeighting:
    """Two elements, two PCs each; 20 theory and 20 practical marks in all."""
    el1 = WeightedElement(
        number="1", title="Identify programming language types",
        pcs=[WeightedPC("1.1", "1", "Programming language types are identified",
                        theory_weight=6, practical_weight=4),
             WeightedPC("1.2", "1", "Language features are described",
                        theory_weight=4, practical_weight=6)],
        stated_theory_total=10, stated_practical_total=10)
    el2 = WeightedElement(
        number="2", title="Apply programming logic",
        pcs=[WeightedPC("2.1", "2", "Algorithms are developed",
                        theory_weight=5, practical_weight=5),
             WeightedPC("2.2", "2", "Programs are debugged",
                        theory_weight=5, practical_weight=5)],
        stated_theory_total=10, stated_practical_total=10)
    return UnitWeighting(unit_title=UNIT_TITLE, cdacc_code=CDACC,
                         isced_code=ISCED, knqf_level="5",
                         elements=[el1, el2],
                         stated_grand_theory=20, stated_grand_practical=20,
                         ratio=(1, 1))


def _written_tool() -> AssessmentTool:
    """CAT 1, written, 30 marks. PC 2.2 is NOT selected - it must show N/A."""
    cat = CatDefinition(cat_id=CAT_1, assessment_type=THEORY, total_marks=30,
                        selected_pcs=["1.1", "1.2", "2.1"], duration_minutes=90)
    allocations = [
        Allocation("1", "Identify programming language types", "1.1",
                   "Programming language types are identified", 6, 10, KNOWLEDGE),
        Allocation("1", "Identify programming language types", "1.2",
                   "Language features are described", 4, 8, UNDERSTANDING),
        Allocation("2", "Apply programming logic", "2.1",
                   "Algorithms are developed", 5, 12, APPLYING),
    ]
    items = [
        Item(number=1, element_number="1", pc_number="1.1", bloom=KNOWLEDGE,
             stem="List four types of programming languages.", marks=10,
             item_format="short_response", scenario_id="S1",
             marking_scheme=[MarkingPoint("Machine and assembly languages", 6),
                             MarkingPoint("High level and 4GL languages", 4)]),
        Item(number=2, element_number="1", pc_number="1.2", bloom=UNDERSTANDING,
             stem="Describe the features of the language in the scenario.",
             marks=8, item_format="extended_response", scenario_id="S1",
             marking_scheme=[MarkingPoint("Strong typing is explained", 5),
                             MarkingPoint("Compilation model is explained", 3)]),
        Item(number=3, element_number="2", pc_number="2.1", bloom=APPLYING,
             stem="Develop an algorithm that sorts a list of marks.", marks=12,
             item_format="short_response", scenario_id="S1",
             marking_scheme=[MarkingPoint("Correct pseudocode structure", 7),
                             MarkingPoint("Correct comparison and swap", 5)]),
    ]
    return AssessmentTool(
        unit_title=UNIT_TITLE, cdacc_code=CDACC, isced_code=ISCED,
        knqf_level="5", programme="ICT Technician", cat=cat,
        allocations=allocations,
        scenarios=[Scenario(id="S1", title="The payroll rewrite",
                            text="A college is rewriting its payroll system.")],
        items=items)


OBSERVATION_TEXT = "Workstation is set up with the correct interpreter"
PRODUCT_TEXT = "Program compiles without errors"


def _practical_tool() -> AssessmentTool:
    """CAT 2, practical, 40 marks: 20 observed, 20 on the product."""
    cat = CatDefinition(cat_id=CAT_2, assessment_type=PRACTICAL, total_marks=40,
                        selected_pcs=["1.1", "2.1"], duration_minutes=180)
    allocations = [
        Allocation("1", "Identify programming language types", "1.1",
                   "Programming language types are identified", 4, 16, APPLYING),
        Allocation("2", "Apply programming logic", "2.1",
                   "Algorithms are developed", 5, 24, ANALYSING),
    ]
    brief = TaskBrief(
        task="Write and run a program that reads twenty marks and ranks them.",
        conditions="Individual work in the computer laboratory.",
        tools_equipment_materials=["Desktop computer", "Python 3 interpreter"],
        safety_requirements=["Observe cable safety", "Use an anti-static mat"],
        time_allowed="3 hours")
    return AssessmentTool(
        unit_title=UNIT_TITLE, cdacc_code=CDACC, isced_code=ISCED,
        knqf_level="5", programme="ICT Technician", cat=cat,
        allocations=allocations, task_brief=brief,
        observation_checklist=[
            ChecklistItem(1, OBSERVATION_TEXT, ["1.1"], 12,
                          sub_parts=["Interpreter version is confirmed"]),
            ChecklistItem(2, "Variables are named to the house convention",
                          ["2.1"], 8)],
        product_checklist=[
            ChecklistItem(1, PRODUCT_TEXT, ["2.1"], 10),
            ChecklistItem(2, "Output ranks all twenty marks correctly",
                          ["2.1"], 10)],
        oral_questions=[
            OralQuestion(1, ["2.1"], "Why did you choose that sorting approach?",
                         ["Names the approach", "Justifies it by data size"])])


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _doc(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def _text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _row_texts(table):
    return [[c.text for c in row.cells] for row in table.rows]


def _int(text: str) -> int:
    """A rendered mark cell: blank means nothing was allocated there."""
    return int(text) if text.strip() else 0


def _all_documents():
    weighting, written = _weighting(), _written_tool()
    return [
        assessment_docs.build_pc_distribution(weighting, written),
        assessment_docs.build_table_of_specifications(written),
        assessment_docs.build_candidates_tool(written),
        assessment_docs.build_assessors_tool(written),
    ]


# --------------------------------------------------------------------------- #
# The header block (all four documents)
# --------------------------------------------------------------------------- #
def test_every_document_header_carries_the_identity():
    for data in _all_documents():
        doc = _doc(data)
        assert doc.tables, "every document carries at least the header table"
        header = _text(_doc(data))
        assert UNIT_TITLE in header
        assert f"ISCED Unit Code: {ISCED}" in header
        assert f"TVET CDACC Unit Code: {CDACC}" in header
        assert "Assessment: CAT 1" in header
        assert "Assessment Type: Written (theory)" in header
        assert "Total Marks: 30" in header


def test_practical_header_names_the_practical_type():
    doc = _doc(assessment_docs.build_candidates_tool(_practical_tool()))
    header = _text(doc)
    assert "Assessment: CAT 2" in header
    assert "Assessment Type: Practical" in header
    assert "Total Marks: 40" in header


# --------------------------------------------------------------------------- #
# 1. PC distribution
# --------------------------------------------------------------------------- #
def _pc_table(weighting=None, tool=None):
    data = assessment_docs.build_pc_distribution(weighting or _weighting(),
                                                 tool or _written_tool())
    doc = _doc(data)
    # tables[0] is the header block; tables[1] is the distribution grid
    return doc, doc.tables[1]


def test_pc_distribution_lists_every_pc_of_the_unit():
    doc, table = _pc_table()
    rows = _row_texts(table)
    assert rows[0] == ["Element", "PC", "Performance Criterion",
                       "Theory Weight", "Practical Weight", "CAT 1 Marks"]
    # 1 header + (2 PCs + sub-total) x 2 elements + grand total
    assert len(rows) == 8
    assert [r[1] for r in rows[1:3]] == ["1.1", "1.2"]
    assert [r[1] for r in rows[4:6]] == ["2.1", "2.2"]


def test_pc_distribution_marks_unselected_pcs_na():
    _, table = _pc_table()
    by_pc = {r[1]: r[5] for r in _row_texts(table)[1:] if r[1]}
    assert by_pc["1.1"] == "10"
    assert by_pc["1.2"] == "8"
    assert by_pc["2.1"] == "12"
    assert by_pc["2.2"] == assessment_docs.NOT_ASSESSED


def test_pc_distribution_subtotal_and_grand_rows_are_populated():
    _, table = _pc_table()
    rows = _row_texts(table)
    first_sub, second_sub, grand = rows[3], rows[6], rows[7]
    assert first_sub[2].startswith("Sub-total: Element 1")
    assert (first_sub[3], first_sub[4], first_sub[5]) == ("10", "10", "18")
    assert second_sub[2].startswith("Sub-total: Element 2")
    assert (second_sub[3], second_sub[4], second_sub[5]) == ("10", "10", "12")
    assert grand[2] == "GRAND TOTAL"
    assert (grand[3], grand[4], grand[5]) == ("20", "20", "30")
    # the CAT column reconciles: 18 + 12 == the CAT total
    assert int(first_sub[5]) + int(second_sub[5]) == int(grand[5])


def test_pc_distribution_element_subtotal_na_when_element_untouched():
    """An element this CAT never reaches sub-totals to N/A, not to zero."""
    tool = _written_tool()
    tool.allocations = [a for a in tool.allocations if a.element_number == "1"]
    _, table = _pc_table(tool=tool)
    rows = _row_texts(table)
    assert rows[6][5] == assessment_docs.NOT_ASSESSED


# --------------------------------------------------------------------------- #
# 2. Table of specifications
# --------------------------------------------------------------------------- #
def _tos_table(tool=None):
    data = assessment_docs.build_table_of_specifications(tool or _written_tool())
    doc = _doc(data)
    return doc, doc.tables[1]


def test_tos_has_a_format_subcolumn_under_every_bloom_level():
    _, table = _tos_table()
    rows = _row_texts(table)
    formats = [label for _, label in assessment_docs.ITEM_FORMAT_COLUMNS]
    # 1 element column + 6 levels x 3 formats + 1 total column
    assert len(rows[0]) == 1 + 6 * len(formats) + 1
    assert rows[0][0] == "Element"
    assert rows[0][-1] == "Total"
    # each level heading spans its format sub-columns (a merged cell repeats)
    assert rows[0][1] == "Knowledge"
    assert rows[0][1 + len(formats)] == "Understanding"
    assert rows[1][1:1 + len(formats)] == formats
    assert rows[1][1:-1] == formats * 6


def test_tos_mcq_columns_are_rendered_but_empty_at_knqf_5():
    """The template demands the column; the format rule forbids the items."""
    _, table = _tos_table()
    rows = _row_texts(table)
    keys = [key for key, _ in assessment_docs.ITEM_FORMAT_COLUMNS]
    mcq_at = [1 + level * len(keys) + keys.index(assessment_docs.MCQ)
              for level in range(6)]
    assert rows[1][mcq_at[0]] == "MCQ"          # rendered
    for row in rows[2:]:                        # data rows + the totals row
        assert all(row[i].strip() == "" for i in mcq_at)
    assert not assessment_docs._mcq_allowed("5")
    assert assessment_docs._mcq_allowed("4")


def test_tos_row_and_column_totals_reconcile_to_the_cat_total():
    tool = _written_tool()
    _, table = _tos_table(tool)
    rows = _row_texts(table)
    data_rows, total_row = rows[2:-1], rows[-1]
    assert len(data_rows) == 2                  # two elements carry allocations
    assert total_row[0] == "TOTAL"

    grand = _int(total_row[-1])
    assert grand == tool.cat.total_marks
    assert sum(_int(r[-1]) for r in data_rows) == grand
    assert sum(_int(v) for v in total_row[1:-1]) == grand
    for row in data_rows:
        assert sum(_int(v) for v in row[1:-1]) == _int(row[-1])


def test_tos_places_marks_in_the_level_and_format_cell():
    _, table = _tos_table()
    rows = _row_texts(table)
    keys = [key for key, _ in assessment_docs.ITEM_FORMAT_COLUMNS]
    def at(level_index, fmt):
        return 1 + level_index * len(keys) + keys.index(fmt)
    # item 1: 10 marks, knowledge, short response
    assert rows[2][at(0, assessment_docs.SHORT_RESPONSE)] == "10"
    # item 2: 8 marks, understanding, extended response
    assert rows[2][at(1, assessment_docs.EXTENDED_RESPONSE)] == "8"
    # item 3: 12 marks, applying, short response, second element
    assert rows[3][at(2, assessment_docs.SHORT_RESPONSE)] == "12"


def test_tos_of_a_practical_falls_back_to_the_allocations():
    tool = _practical_tool()
    _, table = _tos_table(tool)
    rows = _row_texts(table)
    assert _int(rows[-1][-1]) == tool.cat.total_marks


# --------------------------------------------------------------------------- #
# 3. The candidate's tool
# --------------------------------------------------------------------------- #
def test_candidate_written_paper_carries_instructions_items_and_marks():
    tool = _written_tool()
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))
    assert "INSTRUCTIONS TO THE CANDIDATE" in text
    assert "This paper consists of 3 question(s)." in text
    assert "Time allowed: 90 minutes." in text
    for item in tool.items:
        assert item.stem in text
        assert f"({item.marks} marks)" in text


def test_candidate_paper_sets_the_scenario_once_at_the_top():
    tool = _written_tool()
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))
    body = tool.scenarios[0].text
    assert text.count(body) == 1
    assert "Scenario 1: The payroll rewrite" in text


def test_a_single_scenario_is_pointed_at_once_not_on_every_question():
    """The instruction says the paper is about it, so repeating "(Refer to
    Scenario 1)" under all three questions is noise on the page."""
    text = _text(_doc(assessment_docs.build_candidates_tool(_written_tool())))

    assert "(Refer to Scenario 1)" not in text
    assert "Read the scenario in Section A" in text


def test_several_scenarios_are_named_on_the_questions_that_use_them():
    """With a choice of situations the candidate has to be told which one."""
    tool = _written_tool()
    tool.scenarios.append(Scenario(id="S2", title="The stock take",
                                   text="A store is counting its spares."))
    tool.items[2].scenario_id = "S2"
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))

    assert "(Refer to Scenario 1)" in text
    assert "(Refer to Scenario 2)" in text


def test_candidate_paper_never_shows_the_marking_scheme():
    tool = _written_tool()
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))
    for item in tool.items:
        for point in item.marking_scheme:
            assert point.text not in text


def test_candidate_practical_brief_is_complete():
    tool = _practical_tool()
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))
    brief = tool.task_brief
    assert brief.task in text
    assert brief.conditions in text
    for line in brief.tools_equipment_materials + brief.safety_requirements:
        assert line in text
    assert brief.time_allowed in text


def test_candidate_practical_brief_hides_the_checklists_and_the_split():
    tool = _practical_tool()
    text = _text(_doc(assessment_docs.build_candidates_tool(tool)))
    for entry in tool.observation_checklist + tool.product_checklist:
        assert entry.text not in text
        for part in entry.sub_parts:
            assert part not in text
    for question in tool.oral_questions:
        assert question.question not in text
    assert "Observation" not in text
    assert "Checklist" not in text and "checklist" not in text
    # no per-item mark split leaks through the brief
    assert "12" not in text and "Item of Evaluation" not in text


# --------------------------------------------------------------------------- #
# 4. The assessor's tool
# --------------------------------------------------------------------------- #
def test_assessor_written_tool_is_the_marking_scheme():
    tool = _written_tool()
    doc = _doc(assessment_docs.build_assessors_tool(tool))
    table = doc.tables[1]
    rows = _row_texts(table)
    assert rows[0] == list(assessment_docs.MARKING_SCHEME_HEADERS)
    points = sum(len(i.marking_scheme) for i in tool.items)
    assert len(rows) == 1 + points + 1          # header + points + TOTAL
    assert rows[-1][1] == "TOTAL"
    assert rows[-1][3] == str(tool.cat.total_marks)
    text = _text(doc)
    for item in tool.items:
        for point in item.marking_scheme:
            assert point.text in text
    # the marking points sum to their item's marks, as supplied
    assert sum(_int(r[3]) for r in rows[1:-1]) == tool.cat.total_marks


def test_assessor_practical_tool_keeps_the_two_checklists_apart():
    tool = _practical_tool()
    doc = _doc(assessment_docs.build_assessors_tool(tool))
    # header block + observation + product + oral questions
    assert len(doc.tables) == 4
    observation, product, oral = doc.tables[1], doc.tables[2], doc.tables[3]
    assert _row_texts(observation)[0] == list(assessment_docs.CHECKLIST_HEADERS)
    assert _row_texts(product)[0] == list(assessment_docs.CHECKLIST_HEADERS)
    assert _row_texts(oral)[0] == list(assessment_docs.ORAL_HEADERS)
    # 1 header + 2 items + sub-total, on each checklist
    assert len(observation.rows) == 4
    assert len(product.rows) == 4
    assert OBSERVATION_TEXT in _text(doc) and PRODUCT_TEXT in _text(doc)
    assert _row_texts(observation)[-1][3] == "20"
    assert _row_texts(product)[-1][3] == "20"


def test_assessor_practical_tool_carries_oral_response_indicators():
    tool = _practical_tool()
    text = _text(_doc(assessment_docs.build_assessors_tool(tool)))
    question = tool.oral_questions[0]
    assert question.question in text
    for indicator in question.response_indicators:
        assert indicator in text


def test_assessor_practical_tool_omits_the_oral_table_when_there_are_none():
    tool = _practical_tool()
    tool.oral_questions = []
    doc = _doc(assessment_docs.build_assessors_tool(tool))
    assert len(doc.tables) == 3


# --------------------------------------------------------------------------- #
# House style
# --------------------------------------------------------------------------- #
def test_documents_use_the_plan_font_and_table_grid():
    for data in _all_documents():
        doc = _doc(data)
        assert doc.styles["Normal"].font.name == assessment_docs.FONT_NAME
        for table in doc.tables:
            assert table.style.name == "Table Grid"


def test_no_char_splitting_in_list_cells():
    doc = _doc(assessment_docs.build_assessors_tool(_practical_tool()))
    cell = doc.tables[1].rows[1].cells[1]       # text + its sub-parts
    paras = [p.text for p in cell.paragraphs if p.text]
    assert len(paras) == 2
    assert all(len(p) > 1 for p in paras)
