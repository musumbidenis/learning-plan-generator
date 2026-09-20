"""Unit detection across document shapes the original detector could not read."""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import unit_index
from models import UnitRef
from pdf_utils import Page


def page(index, lines):
    """Build a Page from (text, x0, top) triples, mimicking extracted words."""
    words = []
    for text, x0, top in lines:
        x = x0
        for tok in text.split():
            words.append({"text": tok, "x0": x, "x1": x + 5 * len(tok),
                          "top": float(top), "bottom": float(top) + 10})
            x += 5 * len(tok) + 5
    txt = "\n".join(l[0] for l in sorted(lines, key=lambda l: l[2]))
    return Page(index=index, text=txt, words=words)


def unit_page(index, title, code, code_first=False):
    header = [("UNIT CODE: " + code, 80, 50), (title, 80, 70)] if code_first \
        else [(title, 80, 50), ("UNIT CODE: " + code, 80, 70)]
    return page(index, header + [
        ("UNIT DESCRIPTION: This unit specifies the competences required to "
         "do the thing.", 80, 90),
        ("ELEMENT", 80, 130), ("PERFORMANCE CRITERIA", 240, 130),
        ("1. Do the thing", 80, 150), ("1.1 The thing is done", 240, 150)])


def roster_page(index, rows, header="Summary of Units of Learning"):
    lines = [(header, 80, 30),
             ("Unit Code Units Title Unit Duration Credit Factor", 80, 45)]
    for i, (code, title, dur) in enumerate(rows):
        lines.append(("%s %s %d 8" % (code, title, dur), 80, 60 + i * 14))
    lines.append(("Sub Total 480 48", 80, 60 + len(rows) * 14))
    return page(index, lines)


# --------------------------------------------------------------------------- #
# Reading the preliminary units table
# --------------------------------------------------------------------------- #
def test_reads_the_units_table_from_the_front_matter():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)])]
    roster = unit_index.read_roster(pages)
    assert [(e.code, e.title) for e in roster] == [
        ("0611 351 01A", "Computer Essentials"),
        ("0611 351 02A", "Computer Operations")]


def test_reads_a_units_table_that_leads_with_a_category_column():
    pages = [page(0, [
        ("UNIT CATEGORY UNIT CODE UNITS NAME DURATION (HOURS)", 80, 30),
        ("CORE 0611 351 01A Computer Essentials 80", 80, 50),
        ("CORE 0611 351 02A Computer Operations 100", 80, 64),
        ("Sub Total Hours 480", 80, 78)])]
    roster = unit_index.read_roster(pages)
    assert [e.title for e in roster] == ["Computer Essentials",
                                         "Computer Operations"]


def test_reads_a_units_table_of_tvet_codes():
    pages = [page(0, [
        ("Summary of Units of Competency", 80, 30),
        ("IT/OS/ICTA/CC/02/5/MA Apply Communication Skills 60", 80, 50),
        ("IT/OS/ICTA/CC/03/5/MA Apply Digital Literacy 40", 80, 64)])]
    assert [e.title for e in unit_index.read_roster(pages)] == [
        "Apply Communication Skills", "Apply Digital Literacy"]


def test_totals_rows_are_not_mistaken_for_units():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)])]
    titles = [e.title.lower() for e in unit_index.read_roster(pages)]
    assert not any("total" in t for t in titles)


def test_a_units_own_page_is_never_read_as_a_roster():
    """One code on the page - a roster lists several."""
    assert unit_index.read_roster(
        [unit_page(0, "COMPUTER ESSENTIALS", "0611 351 01A")]) == []


def test_a_roster_split_across_two_pages_is_merged():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)]),
             roster_page(1, [("0612 351 03A", "Computer Network Setup", 150),
                             ("0714 351 04A", "Computer Repair", 150)])]
    assert len(unit_index.read_roster(pages)) == 4


# --------------------------------------------------------------------------- #
# The shapes that used to yield nothing
# --------------------------------------------------------------------------- #
def test_finds_units_in_a_document_with_only_tvet_codes():
    """The old detector required an ISCED code and found zero here."""
    pages = [unit_page(0, "APPLY COMMUNICATION SKILLS", "IT/OS/ICTA/CC/02/5/MA"),
             unit_page(1, "APPLY DIGITAL LITERACY", "IT/OS/ICTA/CC/03/5/MA")]
    res = unit_index.index_units(pages, "OS")
    assert [r.title for r in res.refs] == ["APPLY COMMUNICATION SKILLS",
                                           "APPLY DIGITAL LITERACY"]
    assert [r.code for r in res.refs] == ["IT/OS/ICTA/CC/02/5/MA",
                                          "IT/OS/ICTA/CC/03/5/MA"]


def test_finds_a_unit_whose_code_is_printed_above_its_title():
    pages = [unit_page(0, "COMPUTER ESSENTIALS", "0611 351 01A", code_first=True)]
    res = unit_index.index_units(pages, "CU")
    assert [r.title for r in res.refs] == ["COMPUTER ESSENTIALS"]


def test_units_get_contiguous_non_overlapping_page_spans():
    pages = [unit_page(i, "UNIT %d" % i, "061%d 351 0%dA" % (i, i))
             for i in range(3)]
    refs = unit_index.index_units(pages, "OS").refs
    assert [(r.start_page, r.end_page) for r in refs] == [(0, 1), (1, 2), (2, 3)]


def test_an_empty_document_yields_nothing_rather_than_raising():
    res = unit_index.index_units([], "OS")
    assert res.refs == [] and res.roster == [] and res.strategy == "none"


# --------------------------------------------------------------------------- #
# Reporting the shortfall
# --------------------------------------------------------------------------- #
def test_a_rostered_unit_missing_from_the_body_is_reported():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)]),
             unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A")]
    res = unit_index.index_units(pages, "CU")
    assert [r.title.upper() for r in res.refs] == ["COMPUTER ESSENTIALS"]
    assert [m.title for m in res.missing] == ["Computer Operations"]


def test_nothing_is_reported_missing_when_every_rostered_unit_is_found():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)]),
             unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A"),
             unit_page(2, "COMPUTER OPERATIONS", "0611 351 02A")]
    res = unit_index.index_units(pages, "CU")
    assert len(res.refs) == 2
    assert res.missing == []


def test_the_roster_page_itself_is_never_returned_as_a_unit():
    pages = [roster_page(0, [("0611 351 01A", "Computer Essentials", 80),
                             ("0611 351 02A", "Computer Operations", 100)]),
             unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A"),
             unit_page(2, "COMPUTER OPERATIONS", "0611 351 02A")]
    assert all(r.start_page != 0
               for r in unit_index.index_units(pages, "CU").refs)


# --------------------------------------------------------------------------- #
# Multi-unit DOCX
# --------------------------------------------------------------------------- #
def test_a_multi_unit_docx_yields_one_page_per_unit(tmp_path):
    """load_docx_pages used to return ONE page, capping the file at one unit."""
    from docx import Document
    from pdf_utils import load_document

    doc = Document()
    titles = ["APPLY COMMUNICATION SKILLS", "APPLY DIGITAL LITERACY",
              "PERFORM COMPUTER OPERATIONS"]
    for i, title in enumerate(titles, start=1):
        doc.add_paragraph(title)
        doc.add_paragraph("ISCED UNIT CODE: 0611 351 0%dA" % i)
        doc.add_paragraph("UNIT DESCRIPTION: This unit specifies the "
                          "competences required to do the thing.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "ELEMENT"
        t.cell(0, 1).text = "PERFORMANCE CRITERIA"
        t.cell(1, 0).text = "1. Do a thing"
        t.cell(1, 1).text = "1.1 The thing is done"
        doc.add_page_break()
    path = str(tmp_path / "three_units.docx")
    doc.save(path)

    pages = load_document(path)
    assert len(pages) >= 3
    assert [r.title for r in unit_index.index_units(pages, "OS").refs] == titles


# --------------------------------------------------------------------------- #
# A unit that starts under its module's units table, on the same page
# --------------------------------------------------------------------------- #
def module_table_then_unit_page(index):
    """The real shape of a CDACC curriculum: 'MODULE 6', the table of the two
    units in it, then the first of those units starting below."""
    return page(index, [
        ("MODULE 6", 80, 20),
        ("UNIT UNIT CATEGORY ISCED UNIT CODE TVET CDACC UNIT CODE "
         "UNIT NAME DURATION (HOURS)", 80, 34),
        ("CORE 0612 551 IT/CU/ICTA/CR/02/6/MA ICT Security 150", 80, 48),
        ("CORE 0613 551 IT/CU/ICTA/CR/03/6/MA Desktop Application 280", 80, 62),
        ("Total Hours 430", 80, 76),
        ("ICT SECURITY", 80, 100),
        ("ISCED UNIT CODE: 0612 551 16A", 80, 114),
        ("TVET CDACC UNIT CODE: IT/CU/ICTA/CR/02/6/MA", 80, 128),
        ("Unit Description", 80, 142),
        ("This unit covers the competencies required to manage ICT security.",
         80, 156),
        ("Learning Outcomes", 80, 170), ("Content", 300, 170),
        ("1. Assess security needs", 80, 184),
        ("1.1 Security threats", 300, 184)])


def test_a_unit_under_its_module_table_is_named_after_itself():
    """This shipped a unit called 'MODULE 6'.

    The title search anchored on the first code-shaped line, which is a row of
    the table above; the unit's own code line is the one carrying a label.
    """
    refs = unit_index.index_units([module_table_then_unit_page(0)], "CU").refs
    assert [r.title for r in refs] == ["ICT SECURITY"]


def test_the_unit_takes_its_own_code_not_the_tables_clipped_one():
    refs = unit_index.index_units([module_table_then_unit_page(0)], "CU").refs
    assert refs[0].code == "IT/CU/ICTA/CR/02/6/MA"
    assert refs[0].isced_code == "0612 551 16A"


def test_a_units_own_code_header_is_not_read_as_a_table_row():
    """A unit page carries both code families, which made it 'a roster page'
    listing a unit named 'TVET CDACC UNIT CODE:'."""
    roster = unit_index.read_roster([module_table_then_unit_page(0)])
    assert all("CODE" not in e.title.upper() for e in roster)
    assert [e.title for e in roster] == ["ICT Security", "Desktop Application"]


# --------------------------------------------------------------------------- #
# Reconciling the roster against what was found
# --------------------------------------------------------------------------- #
def test_a_roster_row_carrying_the_other_code_family_is_not_reported_missing():
    """The two code columns land in one cell: code '0612 451 07A', title
    'IT/CU/ICTA/CR/02/5/MA Network Design and'. The document quotes the TVET
    code, so comparing only the ISCED one cried wolf on a located unit."""
    entry = unit_index.RosterEntry(
        code="0612 451 07A",
        title="IT/CU/ICTA/CR/02/5/MA Network Design and")
    refs = [UnitRef(title="NETWORK DESIGN AND MANAGEMENT",
                    code="IT/CU/ICTA/CR/02/5/MA")]
    assert unit_index._entry_was_found(entry, refs)


def test_a_clipped_roster_code_still_matches_the_unit():
    """Table cells wrap: 'IT/CU/ICTA/CC/01/6/M' is '...MA' with the last
    character on the next line."""
    entry = unit_index.RosterEntry(code="IT/CU/ICTA/CC/01/6/M", title="Discrete")
    refs = [UnitRef(title="DISCRETE MATHEMATICAL CONCEPTS",
                    code="IT/CU/ICTA/CC/01/6/MA")]
    assert unit_index._entry_was_found(entry, refs)


def test_a_clipped_roster_title_still_matches_the_unit():
    entry = unit_index.RosterEntry(code="0714 351 04A",
                                   title="Computer Repair and")
    refs = [UnitRef(title="COMPUTER REPAIR AND MAINTENANCE",
                    isced_code="0714 351 09A")]
    assert unit_index._entry_was_found(entry, refs)


def test_a_unit_the_document_really_lacks_is_still_reported():
    entry = unit_index.RosterEntry(code="0322 551 11A",
                                   title="Manage Organization Records")
    refs = [UnitRef(title="CONDUCT INDEXING AND ABSTRACTING",
                    isced_code="0322 551 10A")]
    assert not unit_index._entry_was_found(entry, refs)


def test_a_short_shared_prefix_is_not_taken_as_a_match():
    """'Com' opens half the units in an ICT curriculum."""
    entry = unit_index.RosterEntry(code="0611 351 09A", title="Com")
    refs = [UnitRef(title="COMPUTER ESSENTIALS", isced_code="0611 351 01A")]
    assert not unit_index._entry_was_found(entry, refs)


# --------------------------------------------------------------------------- #
# Units the strict detector cannot see
# --------------------------------------------------------------------------- #
def mistyped_unit_page(index, title, isced, tvet):
    """A unit page whose ISCED code is typed with a stray space in it.

    Real: the Agripreneurship level 4 standard heads one unit '0811 34 1 03 A'
    where its own units table says '0811 351 03 A'. The shape-based detector
    requires a well-formed ISCED code and so does not see the page at all.
    """
    return page(index, [
        (title, 80, 50),
        ("UNIT CODE: " + isced, 80, 70),
        ("TVET CDACC UNIT CODE: " + tvet, 80, 85),
        ("UNIT DESCRIPTION: This unit specifies the competences required to "
         "do the thing.", 80, 100),
        ("ELEMENT", 80, 130), ("PERFORMANCE CRITERIA", 240, 130),
        ("1. Do the thing", 80, 150), ("1.1 The thing is done", 240, 150)])


def test_a_unit_whose_own_isced_code_is_mistyped_is_still_found():
    pages = [roster_page(0, [
        ("0811 351 01 A", "AG/OS/PN/CR/01/3/MA Establish Agri-Enterprise", 80),
        ("0811 351 03 A",
         "AG/OS/PN/CR/03/3/MA Market Agri-Enterprise Products and", 120)]),
        unit_page(1, "ESTABLISH AGRI-ENTERPRISE", "0811 351 01 A"),
        mistyped_unit_page(2, "MARKET AGRI-ENTERPRISE PRODUCTS AND SERVICES",
                           "0811 34 1 03 A", "AG/OS/PN/CR/03/3/MA")]

    res = unit_index.index_units(pages, "OS")

    assert [r.title for r in res.refs] == [
        "ESTABLISH AGRI-ENTERPRISE",
        "MARKET AGRI-ENTERPRISE PRODUCTS AND SERVICES"]
    assert res.missing == []


def test_a_unit_that_was_missed_does_not_leave_its_pages_to_its_neighbour():
    """The real damage of a missed unit: refs are page RANGES, so the unit
    before it silently takes its pages and the plan is built from both."""
    pages = [roster_page(0, [
        ("0811 351 01 A", "AG/OS/PN/CR/01/3/MA Establish Agri-Enterprise", 80),
        ("0811 351 03 A",
         "AG/OS/PN/CR/03/3/MA Market Agri-Enterprise Products and", 120)]),
        unit_page(1, "ESTABLISH AGRI-ENTERPRISE", "0811 351 01 A"),
        mistyped_unit_page(2, "MARKET AGRI-ENTERPRISE PRODUCTS AND SERVICES",
                           "0811 34 1 03 A", "AG/OS/PN/CR/03/3/MA")]

    res = unit_index.index_units(pages, "OS")

    assert res.refs[0].end_page == 2


def test_a_unit_put_back_is_titled_from_its_own_page():
    """Not from the roster row, which wraps and arrives with the neighbouring
    code column joined onto the front of it."""
    pages = [roster_page(0, [
        ("0811 351 01 A", "AG/OS/PN/CR/01/3/MA Establish Agri-Enterprise", 80),
        ("0811 351 03 A",
         "AG/OS/PN/CR/03/3/MA Market Agri-Enterprise Products and", 120)]),
        unit_page(1, "ESTABLISH AGRI-ENTERPRISE", "0811 351 01 A"),
        mistyped_unit_page(2, "MARKET AGRI-ENTERPRISE PRODUCTS AND SERVICES",
                           "0811 34 1 03 A", "AG/OS/PN/CR/03/3/MA")]

    titles = [r.title for r in unit_index.index_units(pages, "OS").refs]

    assert titles[1] == "MARKET AGRI-ENTERPRISE PRODUCTS AND SERVICES"


# --------------------------------------------------------------------------- #
# Rows that name no unit
# --------------------------------------------------------------------------- #
def test_industrial_attachment_is_not_reported_as_a_missing_unit():
    """Every curriculum's units table lists it; no curriculum carries a unit
    for it, only a paragraph of hours owed. 146 of the 588 rows the app was
    reporting as impossible to locate were these."""
    pages = [roster_page(0, [
        ("0611 351 01A", "Computer Essentials", 80),
        ("0611 351 02A", "Computer Operations", 100),
        ("0811 251 06 A", "AG/CU/PN/CR/06/3/MA Industrial Training", 480)]),
        unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A"),
        unit_page(2, "COMPUTER OPERATIONS", "0611 351 02A")]

    assert unit_index.index_units(pages, "CU").missing == []


def test_a_real_unit_whose_name_carries_industrial_is_still_reported():
    """'Apply Industrial Chemistry' and 'Perform Industrial Automation' are
    units, which is why the attachment rule is anchored at both ends."""
    pages = [roster_page(0, [
        ("0611 351 01A", "Computer Essentials", 80),
        ("0611 351 02A", "Computer Operations", 100),
        ("0541 541 13A", "Apply Industrial Chemistry", 120)]),
        unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A"),
        unit_page(2, "COMPUTER OPERATIONS", "0611 351 02A")]

    res = unit_index.index_units(pages, "CU")

    assert [m.title for m in res.missing] == ["Apply Industrial Chemistry"]


def test_industrial_attachment_is_not_offered_as_a_unit_either():
    """Looking for a unit that does not exist finds the paragraph of hours
    owed, and puts a unit with no learning outcomes in front of the trainer."""
    pages = [roster_page(0, [
        ("0611 351 01A", "Computer Essentials", 80),
        ("0611 351 02A", "Computer Operations", 100),
        ("0811 251 06 A", "AG/CU/PN/CR/06/3/MA Industrial Training", 480)]),
        unit_page(1, "COMPUTER ESSENTIALS", "0611 351 01A"),
        unit_page(2, "COMPUTER OPERATIONS", "0611 351 02A"),
        page(3, [("Industrial Training", 80, 50),
                 ("UNIT DESCRIPTION: The trainee shall undergo industrial "
                  "training for a minimum of 480 hours.", 80, 70)])]

    titles = [r.title for r in unit_index.index_units(pages, "CU").refs]

    assert titles == ["COMPUTER ESSENTIALS", "COMPUTER OPERATIONS"]


def test_a_unit_the_units_table_never_named_is_still_read():
    """The Forex and Securities standard carries nine units and lists eight.
    Trusting the table alone lost the ninth and ran the eighth to the end of
    the document."""
    pages = [roster_page(0, [
        ("0412 454 02A", "BUS/OS/FRX/CR/01/5/MA Trade Currencies and Stocks", 80),
        ("0412 454 03A", "BUS/OS/FRX/CR/02/5/MA Manage Financial Investments", 80)]),
        unit_page(1, "TRADE CURRENCIES AND STOCKS", "0412 454 02A"),
        unit_page(2, "MANAGE FINANCIAL INVESTMENTS", "0412 454 03A"),
        mistyped_unit_page(3, "COMMUNICATE STOCKS FINANCIAL INFORMATION",
                           "0412 45 4 04A", "BUS/OS/FRX/CR/03/5/MA")]

    res = unit_index.index_units(pages, "OS")

    assert [r.title for r in res.refs][-1] == \
        "COMMUNICATE STOCKS FINANCIAL INFORMATION"
    assert res.refs[1].end_page == 3


def test_a_page_repeating_its_own_units_code_does_not_split_that_unit():
    """What makes an unnamed start believable is carrying a DIFFERENT code
    from the unit it would be splitting; a page in the middle of a unit
    repeats that unit's own code."""
    pages = [roster_page(0, [
        ("0412 454 02A", "BUS/OS/FRX/CR/01/5/MA Trade Currencies and Stocks", 80),
        ("0412 454 03A", "BUS/OS/FRX/CR/02/5/MA Manage Financial Investments", 80)]),
        unit_page(1, "TRADE CURRENCIES AND STOCKS", "0412 454 02A"),
        unit_page(2, "MANAGE FINANCIAL INVESTMENTS", "0412 454 03A"),
        page(3, [("Evidence Guide", 80, 50),
                 ("TVET CDACC UNIT CODE: BUS/OS/FRX/CR/02/5/MA", 80, 70),
                 ("UNIT DESCRIPTION continues here with more prose.", 80, 90)])]

    res = unit_index.index_units(pages, "OS")

    assert len(res.refs) == 2


def test_a_malformed_code_line_is_not_taken_for_the_title():
    """'ISCED UNIT CODE: 0611 451 01' - the trailing letter is missing, so the
    line is not code-shaped, and the label rule only catches a label alone on
    its line. Short and upper-case, it scored as well as the real title above
    it and won the tie on being nearer."""
    pages = [page(1, [
        ("APPLY DIGITAL LITERACY", 80, 50),
        ("ISCED UNIT CODE: 0611 451 01", 80, 65),
        ("TVETCDACC UNIT CODE: AG/OS/PN/BC/01/5/MA", 80, 80),
        ("UNIT DESCRIPTION: This unit covers the competencies required to "
         "demonstrate digital literacy.", 80, 95),
        ("ELEMENT", 80, 130), ("PERFORMANCE CRITERIA", 240, 130)])]

    assert unit_index._best_title(pages[0]) == "APPLY DIGITAL LITERACY"
