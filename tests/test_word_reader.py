"""Reading Word documents of every common format, not just .docx."""

import os
import shutil
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import unit_index
import word_reader
from pdf_utils import load_document
from word_reader import UnsupportedDocument

TITLES = ["APPLY COMMUNICATION SKILLS", "APPLY DIGITAL LITERACY",
          "PERFORM COMPUTER OPERATIONS"]


# --------------------------------------------------------------------------- #
# Fixtures: the same three-unit document written in several formats
# --------------------------------------------------------------------------- #
def _build_docx(path):
    from docx import Document
    doc = Document()
    for i, title in enumerate(TITLES, start=1):
        doc.add_paragraph(title)
        doc.add_paragraph("ISCED UNIT CODE: 0611 351 0%dA" % i)
        doc.add_paragraph("UNIT DESCRIPTION: This unit specifies the "
                          "competences required to do the thing.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "ELEMENT"
        t.cell(0, 1).text = "PERFORMANCE CRITERIA"
        t.cell(1, 0).text = "1. Do a thing"
        t.cell(1, 1).text = "1.1 The thing is done"
        doc.add_page_break()
    doc.save(path)
    return path


def _rtf_source():
    parts = [r"{\rtf1\ansi\deff0",
             r"{\fonttbl{\f0 Times New Roman;}}",
             r"{\info{\author Somebody}}"]
    for i, title in enumerate(TITLES, start=1):
        parts.append(r"\pard " + title + r"\par")
        parts.append(r"\pard ISCED UNIT CODE: 0611 351 0%dA\par" % i)
        parts.append(r"\pard UNIT DESCRIPTION: This unit specifies the "
                     r"competences required to do the thing.\par")
        parts.append(r"\trowd ELEMENT\cell PERFORMANCE CRITERIA\cell\row")
        parts.append(r"\trowd 1. Do a thing\cell 1.1 The thing is done\cell\row")
        parts.append(r"\page")
    parts.append("}")
    return "\n".join(parts)


def _build_rtf(path):
    with open(path, "w", encoding="latin-1") as fh:
        fh.write(_rtf_source())
    return path


_ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
    xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
  <office:body><office:text>
    %s
  </office:text></office:body>
</office:document-content>
"""


def _build_odt(path):
    body = []
    for i, title in enumerate(TITLES, start=1):
        body.append("<text:h>%s</text:h>" % title)
        body.append("<text:p>ISCED UNIT CODE: 0611 351 0%dA</text:p>" % i)
        body.append("<text:p>UNIT DESCRIPTION: This unit specifies the "
                    "competences required to do the thing.</text:p>")
        body.append(
            "<table:table><table:table-row>"
            "<table:table-cell><text:p>ELEMENT</text:p></table:table-cell>"
            "<table:table-cell><text:p>PERFORMANCE CRITERIA</text:p>"
            "</table:table-cell></table:table-row>"
            "<table:table-row>"
            "<table:table-cell><text:p>1. Do a thing</text:p></table:table-cell>"
            "<table:table-cell><text:p>1.1 The thing is done</text:p>"
            "</table:table-cell></table:table-row></table:table>")
        body.append('<text:p><text:soft-page-break/></text:p>')
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("content.xml", _ODT_CONTENT % "\n".join(body))
    return path


# --------------------------------------------------------------------------- #
# Format sniffing - by content, never by extension
# --------------------------------------------------------------------------- #
def test_sniffs_each_format_from_its_magic_bytes(tmp_path):
    assert word_reader.sniff_format(_build_docx(str(tmp_path / "a.docx"))) \
        == word_reader.FMT_OOXML
    assert word_reader.sniff_format(_build_odt(str(tmp_path / "a.odt"))) \
        == word_reader.FMT_ODT
    assert word_reader.sniff_format(_build_rtf(str(tmp_path / "a.rtf"))) \
        == word_reader.FMT_RTF


def test_a_docx_renamed_to_doc_still_opens(tmp_path):
    """Renamed files are a common reason a document 'won't open'."""
    src = _build_docx(str(tmp_path / "real.docx"))
    renamed = str(tmp_path / "mislabelled.doc")
    shutil.copyfile(src, renamed)
    assert word_reader.sniff_format(renamed) == word_reader.FMT_OOXML
    assert [r.title for r in
            unit_index.index_units(load_document(renamed), "OS").refs] == TITLES


def test_an_rtf_named_docx_still_opens(tmp_path):
    src = _build_rtf(str(tmp_path / "real.rtf"))
    renamed = str(tmp_path / "mislabelled.docx")
    shutil.copyfile(src, renamed)
    assert word_reader.sniff_format(renamed) == word_reader.FMT_RTF


def test_a_pdf_is_still_recognised(tmp_path):
    path = str(tmp_path / "x.pdf")
    with open(path, "wb") as fh:
        fh.write(b"%PDF-1.4\n%stub\n")
    assert word_reader.sniff_format(path) == word_reader.FMT_PDF


def test_an_unreadable_file_is_rejected_with_a_helpful_message(tmp_path):
    path = str(tmp_path / "notes.txt")
    with open(path, "w") as fh:
        fh.write("just some notes")
    with pytest.raises(UnsupportedDocument) as excinfo:
        word_reader.read_blocks(path)
    assert ".docx" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# Each format reads into the same blocks, and the same units
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("builder,suffix", [
    (_build_docx, ".docx"),
    (_build_rtf, ".rtf"),
    (_build_odt, ".odt"),
])
def test_every_format_yields_the_same_three_units(builder, suffix, tmp_path):
    path = builder(str(tmp_path / ("doc" + suffix)))
    pages = load_document(path)
    refs = unit_index.index_units(pages, "OS").refs
    assert [r.title for r in refs] == TITLES


@pytest.mark.parametrize("builder,suffix", [
    (_build_docx, ".docx"),
    (_build_rtf, ".rtf"),
    (_build_odt, ".odt"),
])
def test_every_format_preserves_table_rows(builder, suffix, tmp_path):
    """Tables carry the Element/PC and curriculum content columns - losing them
    would leave the column-aware parsers with nothing to split."""
    path = builder(str(tmp_path / ("doc" + suffix)))
    rows = [b for b in word_reader.read_blocks(path) if b.kind == word_reader.ROW]
    assert rows, "no table rows survived"
    assert any(len(r.cells) >= 2 and "ELEMENT" in r.cells[0].upper()
               for r in rows)
    assert any("1.1" in r.text for r in rows)


def test_macro_enabled_and_template_packages_read_like_a_docx(tmp_path):
    """python-docx rejects several of these on content type alone; reading
    word/document.xml directly means one code path covers them all."""
    src = _build_docx(str(tmp_path / "base.docx"))
    for suffix in (".docm", ".dotx", ".dotm"):
        variant = str(tmp_path / ("doc" + suffix))
        shutil.copyfile(src, variant)
        refs = unit_index.index_units(load_document(variant), "OS").refs
        assert [r.title for r in refs] == TITLES, suffix


def test_rtf_metadata_groups_are_not_read_as_document_text(tmp_path):
    path = _build_rtf(str(tmp_path / "doc.rtf"))
    text = " ".join(b.text for b in word_reader.read_blocks(path))
    assert "Somebody" not in text
    assert "Times New Roman" not in text


def test_word_extensions_cover_the_formats_we_read():
    for ext in (".docx", ".docm", ".dotx", ".dotm", ".doc", ".rtf", ".odt"):
        assert ext in word_reader.WORD_EXTENSIONS


# --------------------------------------------------------------------------- #
# Legacy .doc
# --------------------------------------------------------------------------- #
def test_a_doc_without_olefile_explains_what_to_do(tmp_path, monkeypatch):
    """The dependency is optional at runtime; the message must be actionable."""
    path = str(tmp_path / "legacy.doc")
    with open(path, "wb") as fh:
        fh.write(word_reader._OLE_MAGIC + b"\x00" * 512)

    real_import = __builtins__["__import__"] if isinstance(__builtins__, dict) \
        else __builtins__.__import__

    def no_olefile(name, *args, **kwargs):
        if name == "olefile":
            raise ImportError("no olefile")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr("builtins.__import__", no_olefile)
    with pytest.raises(UnsupportedDocument) as excinfo:
        word_reader.read_blocks(path)
    message = str(excinfo.value)
    assert "olefile" in message and "docx" in message


def test_a_corrupt_doc_fails_with_a_clear_message(tmp_path):
    path = str(tmp_path / "broken.doc")
    with open(path, "wb") as fh:
        fh.write(word_reader._OLE_MAGIC + b"\x00" * 2048)
    with pytest.raises(Exception) as excinfo:
        word_reader.read_blocks(path)
    assert "doc" in str(excinfo.value).lower()
