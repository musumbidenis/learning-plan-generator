"""The trainer's own notes, slides and recordings - the first source of a question.

Everything else this module sits beside is a guess at what was taught. The
curriculum lists topic headings, the occupational standard lists criteria, and
`assessment_knowledge` goes and reads a public encyclopaedia about them. All
three are approximations of the one thing that actually happened: a trainer
stood in a room and taught from something.

That something is what a candidate can fairly be asked about, so when it is
supplied it outranks the rest. A question comes from the attached resources;
the curriculum says which topics are in scope and the occupational standard
says what the marks are for; and only where the resources are silent does
anything else fill the gap.

WHAT CAN BE READ
  .pdf                     pages, via the project's own loader
  .docx .doc .odt .rtf     paragraphs and table rows, via `word_reader`
  .pptx                    one chunk per slide - title, body, speaker notes
  .txt .md                 as written, split on headings or blank lines
  .mp3 .m4a .wav .mp4 ...  transcribed by Groq's Whisper

A recording is transcribed rather than watched. There is no ffmpeg here, so
the file goes to the transcription endpoint as it stands, and that endpoint
takes a video container and pulls the audio out itself. What comes back is
what was SAID. Slides shown silently, a diagram drawn on a board, a
demonstration performed without narration - none of that is in the transcript,
and a paper written from one should be read with that in mind. Attach the
slides as well and the two cover each other.

WHY IT IS CHUNKED AND SCORED RATHER THAN SENT WHOLE
A unit's notes run to tens of thousands of characters and the whole prompt has
to fit inside about eight thousand tokens. So a resource is cut into pieces at
its own seams - slides, headings, paragraphs - each piece is scored against
the topics this CAT actually assesses, and the best are taken round-robin
ACROSS THE TOPICS so that every assessed topic gets material before any topic
gets a second piece. Taking the best pieces overall instead would hand the
whole budget to whichever topic the notes happen to dwell on.

Nothing here raises. A file that cannot be read costs the paper that file and
not its generation, and says so.
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

import runlog
from assessment_models import ResourceChunk

# How much of the prompt the resources may take. They are the primary source,
# so they get the largest single share of what is left after the standing
# instructions - more than the teaching notes, which now only fill gaps.
RESOURCE_BLOCK_CHARS = 6000

# One piece of a resource. Big enough to carry an idea, small enough that a
# budget buys several from different topics.
CHUNK_CHARS = 900
MIN_CHUNK_CHARS = 60

# A chunk has to share this much of a topic's vocabulary to count as being
# about that topic. Low on purpose: a topic is three or four stemmed words, so
# a chunk matching one of three is already on subject - "Vulnerability
# scanning tools" against a page naming Nessus and OpenVAS scores exactly a
# third, and that page is the best answer in the file.
MIN_TOPIC_OVERLAP = 0.25

# Recordings. Groq's endpoint takes the container and extracts the audio, so
# no ffmpeg is needed - but it caps the upload, and a lecture video passes that
# long before a lecture audio does.
TRANSCRIBE_MODEL = "whisper-large-v3-turbo"
TRANSCRIBE_URL = "https://api.groq.com/openai/v1/audio/transcriptions"
TRANSCRIBE_LIMIT_MB = 24
TRANSCRIBE_TIMEOUT = 300

TEXT_EXTENSIONS = (".txt", ".md", ".markdown", ".csv")
SLIDE_EXTENSIONS = (".pptx", ".potx", ".ppsx")
MEDIA_EXTENSIONS = (".mp3", ".m4a", ".wav", ".webm", ".mp4", ".mpeg", ".mpga",
                    ".ogg", ".flac")

_RE_WORD = re.compile(r"[a-z]{4,}")
_RE_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$")
_RE_BLANK = re.compile(r"\n\s*\n")

_STOP = frozenset("""
the and for with that this from are was were have has had you your they them
their which what when where who how why all any one two three four five six
into out over under between about because during than then some many much
used using use also such other more most only very well often shall will
""".split())


# The shared shape lives in `assessment_models`, with the rest of the
# contract the documents and the prompt are built from.
Chunk = ResourceChunk


@dataclass
class Resource:
    """One attached file, read into chunks."""
    name: str
    kind: str = ""
    chunks: List[Chunk] = field(default_factory=list)
    note: str = ""                   # why it is empty, when it is

    @property
    def ok(self) -> bool:
        return bool(self.chunks)


# --------------------------------------------------------------------------- #
# Reading each kind
# --------------------------------------------------------------------------- #
def _stem(word: str) -> str:
    """A word reduced far enough that its forms meet.

    Not linguistics - a suffix strip and a doubled-letter fix, which is all
    that is needed here and needs no dependency. It exists because the first
    version compared whole words and quietly failed: a curriculum topic saying
    "Vulnerability SCANNING tools" scored zero against a page of notes saying
    "the SCANNERS used in the workshop", and the one resource that answered
    the question was dropped from the prompt.

        scanning, scanners, scan  ->  scan
        vulnerabilities           ->  vulnerability
        threats                   ->  threat
    """
    for suffix, shortest in (("ies", 4), ("ing", 6), ("ers", 5), ("er", 5),
                             ("ed", 5), ("es", 5), ("s", 5)):
        if word.endswith(suffix) and len(word) >= shortest:
            word = word[:-len(suffix)] + ("y" if suffix == "ies" else "")
            break
    if len(word) > 3 and word[-1] == word[-2]:
        word = word[:-1]              # scann -> scan, plann -> plan
    return word


def _tokens(text: str) -> set:
    return {_stem(w) for w in _RE_WORD.findall((text or "").lower())
            if w not in _STOP}


def _split(text: str, where: str, source: str,
           heading: str = "") -> List[Chunk]:
    """One passage as however many chunks of about CHUNK_CHARS it needs.

    Paragraphs are kept whole where they fit, because a paragraph is the
    smallest thing that still makes sense on its own; only a paragraph longer
    than the budget is cut, and then at a sentence.
    """
    out: List[Chunk] = []
    current = ""
    for para in _RE_BLANK.split(text or ""):
        para = " ".join(para.split())
        if not para:
            continue
        while len(para) > CHUNK_CHARS:
            cut = para.rfind(". ", 0, CHUNK_CHARS)
            cut = cut + 1 if cut > MIN_CHUNK_CHARS else CHUNK_CHARS
            out.append(Chunk(source, where, heading, para[:cut].strip()))
            para = para[cut:].strip()
        if len(current) + len(para) + 1 > CHUNK_CHARS and current:
            out.append(Chunk(source, where, heading, current.strip()))
            current = ""
        current = f"{current} {para}".strip()
    if len(current) >= MIN_CHUNK_CHARS:
        out.append(Chunk(source, where, heading, current.strip()))
    return out


def _read_text(name: str, data: bytes) -> List[Chunk]:
    """Plain text or markdown, cut at its own headings where it has them."""
    body = data.decode("utf-8", errors="replace")
    sections: List[Tuple[str, str]] = []
    heading, buffer = "", []
    for line in body.splitlines():
        found = _RE_HEADING.match(line)
        if found:
            if buffer:
                sections.append((heading, "\n".join(buffer)))
            heading, buffer = found.group(1).strip(), []
        else:
            buffer.append(line)
    if buffer:
        sections.append((heading, "\n".join(buffer)))
    out: List[Chunk] = []
    for head, text in sections:
        out.extend(_split(text, "", name, head))
    return out


def _read_pdf(name: str, data: bytes) -> List[Chunk]:
    import pdf_utils
    import tempfile
    handle, path = tempfile.mkstemp(suffix=".pdf")
    try:
        with os.fdopen(handle, "wb") as fh:
            fh.write(data)
        pages = pdf_utils.load_pdf_pages(path)
        out: List[Chunk] = []
        for page in pages:
            out.extend(_split(page.text, f"page {page.index + 1}", name))
        return out
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


def _read_word(name: str, data: bytes) -> List[Chunk]:
    """Any word-processor format the project already reads.

    A table row becomes one line rather than one chunk: a table of tool names
    is a single idea, and a chunk per row would spend the whole budget on it.
    """
    import tempfile
    import word_reader
    suffix = os.path.splitext(name)[1] or ".docx"
    handle, path = tempfile.mkstemp(suffix=suffix)
    try:
        with os.fdopen(handle, "wb") as fh:
            fh.write(data)
        blocks = word_reader.read_blocks(path)
    finally:
        try:
            os.remove(path)
        except OSError:
            pass
    lines = []
    for block in blocks:
        if block.kind == word_reader.ROW:
            lines.append(" | ".join(c for c in block.cells if c))
        elif block.kind == word_reader.BREAK:
            lines.append("")
        else:
            lines.append(block.text)
    return _split("\n\n".join(lines), "", name)


def _read_slides(name: str, data: bytes) -> List[Chunk]:
    """One chunk per slide - a deck is already chunked, and by its author.

    Speaker notes are taken too. They are usually the half a trainer actually
    says out loud, and a bullet reading "Types of malware" with a note under it
    explaining each one is a note worth far more than the bullet.
    """
    from pptx import Presentation
    deck = Presentation(io.BytesIO(data))
    out: List[Chunk] = []
    for number, slide in enumerate(deck.slides, start=1):
        title, body = "", []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(p.text.strip()
                            for p in shape.text_frame.paragraphs
                            if p.text.strip())
            if not text:
                continue
            if shape == getattr(slide.shapes, "title", None) and not title:
                title = text
            else:
                body.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
        whole = " ".join(x for x in (" ".join(body), notes) if x).strip()
        if len(whole) < MIN_CHUNK_CHARS and not title:
            continue
        out.extend(_split(whole or title, f"slide {number}", name, title))
    return out


def _read_media(name: str, data: bytes, api_key: str = "") -> List[Chunk]:
    """A recording, as the words that were spoken in it.

    Sent to the transcription endpoint as it stands. That endpoint accepts a
    video container and takes the audio out itself, which is what makes this
    work at all on a machine with no ffmpeg.
    """
    import requests
    from ai_client import load_api_key
    key = api_key or load_api_key()
    if not key:
        raise RuntimeError("no GROQ_API_KEY, so a recording cannot be "
                           "transcribed")
    size_mb = len(data) / (1024 * 1024)
    if size_mb > TRANSCRIBE_LIMIT_MB:
        raise RuntimeError(
            f"{size_mb:.0f} MB is over the {TRANSCRIBE_LIMIT_MB} MB the "
            f"transcription service accepts - export the audio on its own, or "
            f"attach the slides instead")
    resp = requests.post(
        TRANSCRIBE_URL, headers={"Authorization": f"Bearer {key}"},
        files={"file": (os.path.basename(name), data)},
        data={"model": TRANSCRIBE_MODEL, "response_format": "text"},
        timeout=TRANSCRIBE_TIMEOUT)
    if resp.status_code != 200:
        raise RuntimeError(f"the transcription service answered "
                           f"{resp.status_code}: {resp.text[:120]}")
    spoken = " ".join((resp.text or "").split())
    if len(spoken) < MIN_CHUNK_CHARS:
        raise RuntimeError("nothing was said in it that could be transcribed")
    return _split(spoken, "spoken", name)


_READERS = (
    (TEXT_EXTENSIONS, "notes", _read_text),
    ((".pdf",), "notes", _read_pdf),
    (SLIDE_EXTENSIONS, "slides", _read_slides),
    (MEDIA_EXTENSIONS, "recording", _read_media),
)


def read(name: str, data: bytes, api_key: str = "") -> Resource:
    """One attached file, read into chunks. Never raises.

    A file that cannot be read is returned empty with the reason on it, so the
    trainer sees WHICH resource was not used and why, rather than wondering
    why the paper ignored it.
    """
    lower = (name or "").lower()
    for extensions, kind, reader in _READERS:
        if not lower.endswith(extensions):
            continue
        try:
            chunks = (reader(name, data, api_key)
                      if reader is _read_media else reader(name, data))
        except Exception as e:                            # noqa: BLE001
            runlog.warn(f"Resources: '{name}' could not be read: {e}")
            return Resource(name=name, kind=kind, note=str(e)[:200])
        if not chunks:
            return Resource(name=name, kind=kind,
                            note="no readable text in it")
        runlog.log(f"Resources: '{name}' read as {len(chunks)} piece(s)")
        return Resource(name=name, kind=kind, chunks=chunks)

    from word_reader import WORD_EXTENSIONS
    if lower.endswith(WORD_EXTENSIONS):
        try:
            chunks = _read_word(name, data)
        except Exception as e:                            # noqa: BLE001
            runlog.warn(f"Resources: '{name}' could not be read: {e}")
            return Resource(name=name, kind="notes", note=str(e)[:200])
        return Resource(name=name, kind="notes", chunks=chunks,
                        note="" if chunks else "no readable text in it")

    return Resource(name=name, kind="",
                    note="not a format this reads - attach a PDF, Word "
                         "document, PowerPoint, text file or recording")


# --------------------------------------------------------------------------- #
# Choosing what the prompt carries
# --------------------------------------------------------------------------- #
def topic_keys(content) -> List[Tuple[str, str, set]]:
    """[(topic number, topic title, its vocabulary)] for the assessed topics."""
    out: List[Tuple[str, str, set]] = []
    for block in content or []:
        for topic in block.topics:
            words = _tokens(" ".join([topic.title] + list(topic.key_points)))
            if words:
                out.append((topic.number, topic.title, words))
    return out


def _relevance(chunk: Chunk, wanted: set) -> float:
    """How much of a topic's vocabulary this chunk contains."""
    if not wanted:
        return 0.0
    return len(_tokens(chunk.text) & wanted) / len(wanted)


def select(resources: Sequence[Resource], content,
           budget: int = RESOURCE_BLOCK_CHARS) -> List[Chunk]:
    """The pieces the prompt carries, spread across the assessed topics.

    Round-robin across topics, best piece first, until the budget runs out.
    Taking the best pieces overall would give the whole budget to whichever
    topic the notes dwell on, and the items on every other topic would be
    written from nothing - the same fault, and the same fix, as in
    `assessment_knowledge.key_points`.

    With no curriculum to match against, the resources are taken in the order
    they were attached: the trainer's own order is better than none.
    """
    pool: List[Chunk] = [c for r in resources for c in r.chunks]
    if not pool:
        return []

    topics = topic_keys(content)
    if not topics:
        return _take(pool, budget)

    ranked: List[List[Chunk]] = []
    for _number, _title, wanted in topics:
        scored = [(_relevance(c, wanted), i, c) for i, c in enumerate(pool)]
        keep = [c for score, _i, c in sorted(scored, key=lambda s: (-s[0], s[1]))
                if score >= MIN_TOPIC_OVERLAP]
        if keep:
            ranked.append(keep)

    if not ranked:
        # Nothing matched. The resources are still the trainer's, so they are
        # used rather than dropped - a paper written from off-topic notes is
        # visible to the trainer, and one written from nothing is not.
        runlog.warn("Resources: none of the attached material matched the "
                    "assessed topics; using it in the order attached")
        return _take(pool, budget)

    out: List[Chunk] = []
    seen = set()
    depth = 0
    spent = 0
    while depth < max(len(r) for r in ranked):
        progressed = False
        for queue in ranked:
            if depth >= len(queue):
                continue
            chunk = queue[depth]
            if id(chunk) in seen:
                continue
            if spent + len(chunk.text) > budget and out:
                return out
            seen.add(id(chunk))
            out.append(chunk)
            spent += len(chunk.text)
            progressed = True
        if not progressed:
            break
        depth += 1
    return out


def fit(chunks: Sequence[Chunk], budget: int) -> List[Chunk]:
    """The already-chosen pieces, trimmed again to a smaller budget.

    `select` picks WHICH pieces, once, against the assessed topics. This is
    the second trim, used only when a prompt has to be made smaller to be
    sent at all - and it keeps the order `select` settled on, which is
    round-robin across topics, so shrinking the budget costs the last piece of
    each topic rather than every piece of the last topic.
    """
    return _take(chunks, budget) if budget > 0 else list(chunks)


def _take(pool: Sequence[Chunk], budget: int) -> List[Chunk]:
    out, spent = [], 0
    for chunk in pool:
        if spent + len(chunk.text) > budget and out:
            break
        out.append(chunk)
        spent += len(chunk.text)
    return out


def covered(chunks: Sequence[Chunk], key_point: str) -> bool:
    """Whether the supplied material already says something about this point.

    Used to decide where the gap-filling lookup is still needed. A key point
    the trainer's own notes cover does not need an encyclopaedia.
    """
    wanted = _tokens(key_point)
    if not wanted:
        return False
    return any(len(_tokens(c.text) & wanted) / len(wanted) >= MIN_TOPIC_OVERLAP
               for c in chunks)


def render(chunks: Sequence[Chunk]) -> str:
    """The chosen pieces as the model reads them, or '' when there are none."""
    if not chunks:
        return ""
    by_source: Dict[str, List[Chunk]] = {}
    for chunk in chunks:
        by_source.setdefault(chunk.source, []).append(chunk)
    blocks: List[str] = []
    for source, pieces in by_source.items():
        lines = [f"FROM: {source}"]
        for piece in pieces:
            where = f" [{piece.where}]" if piece.where else ""
            head = f" {piece.heading}" if piece.heading else ""
            lines.append(f"  -{where}{head}")
            lines.append(f"    {piece.text}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def summarise(resources: Sequence[Resource], chosen: Sequence[Chunk]) -> str:
    """One line for the progress log and the UI."""
    usable = [r for r in resources if r.ok]
    if not usable:
        return "no resources attached"
    pieces = sum(len(r.chunks) for r in usable)
    return (f"{len(usable)} resource(s), {pieces} piece(s) read, "
            f"{len(chosen)} used in the prompt")
